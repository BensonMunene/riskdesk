import csv
import io

from django.contrib.auth.decorators import login_required
from django.http import HttpResponse
from django.shortcuts import redirect, render

from web.services import jsonable

from .common import portfolio_context


@login_required
def risk_report(request, pk):
    ctx = portfolio_context(pk, request)
    if ctx["empty"]:
        return redirect("positions", pk=pk)
    a = ctx["a"]
    stress = sorted([d for d in a.stress_historical + a.stress_hypothetical if d.get("available")], key=lambda d: d["total_pnl"])
    ctx.update({
        "rc": a.rc.head(15), "sector_rc": a.sector_rc, "sector_weights": a.sector_weights,
        "exposures": a.factor_decomposition["exposures"], "fd": a.factor_decomposition,
        "var_rows": [v.as_dict() for v in a.var_table], "perf": a.performance,
        "stress": stress[:8], "insights": a.insights(ctx["limit_results"]), "concentration": a.concentration,
        "liquidity_profile": a.liquidity_profile, "drawdowns": a.drawdowns.head(3), "print_mode": request.GET.get("print") == "1",
    })
    return render(request, "web/report.html", ctx)


@login_required
def export_positions_csv(request, pk):
    ctx = portfolio_context(pk, request)
    a = ctx["a"]
    resp = HttpResponse(content_type="text/csv")
    resp["Content-Disposition"] = f'attachment; filename="{ctx["portfolio"].name.replace(" ", "_")}_positions.csv"'
    w = csv.writer(resp)
    w.writerow(["ticker", "name", "sector", "asset_class", "quantity", "price", "market_value", "weight", "risk_share", "standalone_vol", "mcr"])
    rc = a.rc if not ctx["empty"] else None
    for t, row in a.positions.iterrows():
        px = float(a.last_prices[t])
        w.writerow([t, row["name"], row["sector"], row["asset_class"], row["quantity"], round(px, 4),
                    round(row["quantity"] * px, 2), round(float(a.weights[t]), 6),
                    round(float(rc["pct"].get(t, 0.0)), 6) if rc is not None else "",
                    round(float(rc["standalone_vol"].get(t, 0.0)), 6) if rc is not None else "",
                    round(float(rc["mcr"].get(t, 0.0)), 6) if rc is not None else ""])
    return resp


@login_required
def export_risk_xlsx(request, pk):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    ctx = portfolio_context(pk, request)
    if ctx["empty"]:
        return redirect("positions", pk=pk)
    a = ctx["a"]
    wb = Workbook()

    def sheet(title, header, rows):
        ws = wb.create_sheet(title) if wb.active.title != "Sheet" else wb.active
        ws.title = title
        ws.append(header)
        for c in ws[1]:
            c.font = Font(bold=True)
        for r in rows:
            ws.append([("" if v is None else v) for v in r])
        for i, _ in enumerate(header, 1):
            ws.column_dimensions[get_column_letter(i)].width = 18
        return ws

    s = jsonable(ctx["summary"])
    sheet("Summary", ["metric", "value"], [(k, v) for k, v in s.items()])
    rc = a.rc
    sheet("Positions", ["ticker", "name", "sector", "quantity", "price", "market_value", "weight", "risk_share", "ccr", "mcr", "standalone_vol"],
          [(t, r["name"], r["sector"], float(a.positions.loc[t, "quantity"]), float(a.last_prices[t]), float(r["market_value"]),
            float(r["weight"]), float(r["pct"]), float(r["ccr"]), float(r["mcr"]), float(r["standalone_vol"])) for t, r in rc.iterrows()])
    sheet("Sectors", ["sector", "net", "gross", "risk_share", "n"],
          [(sname, float(r["weight"]), float(r["gross"]), float(r["pct"]), int(r["n"])) for sname, r in a.sector_rc.iterrows()])
    ex = a.factor_decomposition["exposures"]
    sheet("Factors", ["factor", "beta", "factor_vol", "pct_of_variance"],
          [(f, float(r["beta"]), float(r["factor_vol"]), float(r["pct_of_total_var"])) for f, r in ex.iterrows()])
    sheet("VaR", ["method", "confidence", "horizon_days", "var", "cvar", "var_dollar", "cvar_dollar"],
          [(v.method, v.confidence, v.horizon_days, v.var, v.cvar, v.var * a.nav, v.cvar * a.nav) for v in a.var_table])
    sheet("Stress", ["scenario", "kind", "pnl_pct", "pnl_dollar", "benchmark"],
          [(d["name"], d["kind"], d["total_pnl"], d["total_dollar"], d.get("benchmark_return"))
           for d in a.stress_historical + a.stress_hypothetical if d.get("available")])
    sheet("Limits", ["limit", "value", "threshold", "status", "utilization"],
          [(l["label"], l["display_value"], l["display_threshold"], l["status"], l.get("utilization")) for l in ctx["limit_results"]])
    sheet("Insights", ["severity", "category", "title", "detail", "action"],
          [(i.severity, i.category, i.title, i.detail, i.action) for i in a.insights(ctx["limit_results"])])
    corr = a.corr
    ws = wb.create_sheet("Correlation")
    ws.append([""] + list(corr.columns))
    for t in corr.index:
        ws.append([t] + [round(float(v), 3) for v in corr.loc[t]])
    buf = io.BytesIO()
    wb.save(buf)
    resp = HttpResponse(buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    resp["Content-Disposition"] = f'attachment; filename="{ctx["portfolio"].name.replace(" ", "_")}_risk_pack.xlsx"'
    return resp


@login_required
def api_docs(request):
    return render(request, "web/api_docs.html")
